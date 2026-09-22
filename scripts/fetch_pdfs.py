#!/usr/bin/env python3
"""
Download PDFs for DOIs listed in chatbots/*.md into pdfs/<chatbot_folder>/,
or resolve author–year citations from an XLSX column (default header ``Author (date)``)
via Crossref/OpenAlex, optionally re-ranked using topic words from extra columns
(``--context-column``) or ``--keyword-boost``, write ``pdfs/<stem>_download_summary.md``
and ``pdfs/<stem>_download_summary.csv``, then download PDFs.

Environment (optional):
  AI4RESEARCH_SCHOLARLY=1 — with --xlsx/--cite (or together with --scholarly), after Crossref fails try Google Scholar via optional scholarly package.
  AI4RESEARCH_UNPAYWALL_EMAIL — override mailto used for Unpaywall API and OpenAlex User-Agent (default: in-script EMAIL).
  AI4RESEARCH_CORE_API_KEY — CORE API v3 bearer token when set (helps quota).
  AI4RESEARCH_COOKIE_FILE — Netscape/Mozilla cookies.txt for institutional PDF access.
  AI4RESEARCH_HTTP_COOKIE — raw Cookie header (alternative to file).
  AI4RESEARCH_USE_CURL_CFFI — (legacy) curl-cffi is now used automatically when installed; this variable has no effect.
  AI4RESEARCH_PLAYWRIGHT=1 — last-resort PDF fetch via headless Chromium (requires: pip install playwright && playwright install).
  AI4RESEARCH_CROSSREF_DEBUG=1 — log Crossref request URL and every hit kept after the first-author filter, with scores (stderr).
  AI4RESEARCH_WEB_HTML_RESOLVE=1 — after Crossref/OpenAlex (and optional Scholarly), fetch DuckDuckGo Lite HTML and accept the first DOI whose Crossref metadata passes the first-author check (same as ``--web-html-resolve``).

Resolution order: Unpaywall -> Semantic Scholar -> OpenAlex -> Crossref ->
 Europe PMC (DOI) -> CORE -> PMC concrete PDFs (Europe PMC by PMCID) ->
 arXiv/PMC guesses -> landing HTML parsing.
PDF fetch order per URL: curl_cffi (Chrome TLS, auto when installed) -> urllib -> curl -> optional Playwright.
Never commit real cookie files.

CLI: python fetch_pdfs.py [--doi DOI] [--out PATH] [--dry-run]
      python fetch_pdfs.py --cite AUTHOR_YEAR [--out PATH] [--resolve-only]
          [--keyword-boost WORDS] [--scholarly] [--search-urls] [--web-html-resolve]
      python fetch_pdfs.py --xlsx PATH [--xlsx-out-dir DIR]
                                [--citation-column NAME] [--context-column NAME ...]
          [--scholarly] [--search-urls] [--web-html-resolve] [--resolve-only]
      (no args: scan chatbots/*.md as before)

  XLSX mode requires: pip install openpyxl

  Google Scholar fallback (optional): pip install scholarly
  Enable with --scholarly (with --xlsx/--cite) or AI4RESEARCH_SCHOLARLY=1 after Crossref fails.
  Scholarly may rate-limit or block; see https://github.com/scholarly-python-package/scholarly
  ``--search-urls`` prints browser-ready Google/Scholar/DuckDuckGo queries when resolution fails.
  ``--web-html-resolve`` (or AI4RESEARCH_WEB_HTML_RESOLVE=1) tries DuckDuckGo Lite HTML after APIs fail.
"""
from __future__ import annotations

import argparse
import csv
import html as html_module
import json
import os
import random
import re
import shutil
import subprocess
import tempfile
import time
import sys
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from http.cookiejar import MozillaCookieJar
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CHATBOTS = ROOT / "chatbots"
PDFS = ROOT / "pdfs"
EMAIL = "guido.biele@fhi.no"


def _unpaywall_email() -> str:
    v = os.environ.get("AI4RESEARCH_UNPAYWALL_EMAIL", "").strip()
    return v if v else EMAIL


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36"
)

CTX = __import__("ssl").create_default_context()

# Max retries for transient HTTP errors (APIs + PDF hosts)
_MAX_RETRIES = 4
_RETRYABLE_STATUS = {429, 500, 502, 503, 504}
_MIN_PDF_BYTES = 800
_MAX_PDF_BYTES = 120 * 1024 * 1024

# Lightweight in-run throttling for API sources to reduce 429s.
_SOURCE_MIN_INTERVAL_S = {
    "unpaywall": 0.10,
    "semanticscholar": 0.50,
    "openalex": 0.10,
    "crossref": 0.10,
    "europepmc": 0.12,
    "core": 0.25,
}
_source_next_allowed_at: dict[str, float] = {}

# Per-source DOI caches (run-local): avoid repeating same lookup.
_lookup_json_cache: dict[tuple[str, str], dict | None] = {}
_lookup_urls_cache: dict[tuple[str, str], tuple[str, ...]] = {}


def _body_snippet(data: bytes, max_len: int = 220) -> str:
    return data[:max_len].decode("utf-8", errors="replace").replace("\n", " ")


def _rate_limit_source(source: str) -> None:
    now = time.time()
    nxt = _source_next_allowed_at.get(source, 0.0)
    if nxt > now:
        time.sleep(nxt - now)
    wait = _SOURCE_MIN_INTERVAL_S.get(source, 0.0)
    _source_next_allowed_at[source] = time.time() + wait


def _validate_pdf_bytes(data: bytes) -> tuple[bool, str]:
    if not data:
        return False, "empty response"
    n = len(data)
    if n < _MIN_PDF_BYTES:
        return False, f"too small ({n} bytes)"
    if n > _MAX_PDF_BYTES:
        return False, f"too large ({n} bytes)"
    if data[:4] != b"%PDF":
        head = data[:2048].lower()
        if b"<html" in head or b"<!doctype" in head:
            return False, "html instead of pdf"
        return False, "missing %PDF header"
    first = data[:5000].lower()
    for marker in (
        b"access denied",
        b"subscription required",
        b"please login",
        b"sign in",
        b"unauthorized",
        b"403 forbidden",
    ):
        if marker in first:
            return False, f"access-error content ({marker.decode(errors='ignore')})"
    return True, ""


def _decompress_if_needed(data: bytes, resp) -> bytes:
    """Decompress gzip/deflate response bodies (urllib does not auto-decompress)."""
    try:
        enc = (resp.headers.get("Content-Encoding") or "").lower()
    except Exception:
        return data
    if "gzip" in enc:
        import gzip as _gzip
        try:
            return _gzip.decompress(data)
        except Exception:
            pass
    elif "deflate" in enc:
        import zlib as _zlib
        try:
            return _zlib.decompress(data)
        except Exception:
            try:
                return _zlib.decompress(data, -15)
            except Exception:
                pass
    return data


def openalex_user_agent() -> str:
    """OpenAlex requests a mailto in the User-Agent string (polite pool)."""
    return f"{USER_AGENT} mailto:{_unpaywall_email()}"


# DOI path may contain parentheses, e.g. 10.1016/S1573-4412(07)06070-9
DOI_RE = re.compile(r"https?://doi\.org/(10\.\d{4,9}/[^\s\]]+)", re.I)

CURL = shutil.which("curl") or shutil.which("curl.exe")

_cookie_opener: urllib.request.OpenerDirector | None = None
_cookie_opener_source: str | None = None
_semantic_scholar_backoff_until: float = 0.0


def _cookie_file_path() -> str | None:
    raw = os.environ.get("AI4RESEARCH_COOKIE_FILE", "").strip()
    if not raw:
        return None
    path = Path(raw).expanduser()
    return str(path.resolve()) if path.is_file() else None


def _cookie_header_value() -> str | None:
    h = os.environ.get("AI4RESEARCH_HTTP_COOKIE", "").strip()
    return h or None


def _get_cookie_opener() -> urllib.request.OpenerDirector | None:
    global _cookie_opener, _cookie_opener_source
    fp = _cookie_file_path()
    if fp is None:
        return None
    if _cookie_opener is not None and _cookie_opener_source == fp:
        return _cookie_opener
    try:
        cj = MozillaCookieJar(fp)
        cj.load(ignore_discard=True, ignore_expires=True)
    except OSError as e:
        print(f"AI4RESEARCH_COOKIE_FILE: could not load {fp}: {e}", file=sys.stderr)
        return None
    except Exception as e:
        print(f"AI4RESEARCH_COOKIE_FILE: invalid or unsupported jar {fp}: {e}", file=sys.stderr)
        return None
    https_handler = urllib.request.HTTPSHandler(context=CTX)
    _cookie_opener = urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(cj),
        https_handler,
    )
    _cookie_opener_source = fp
    return _cookie_opener


def _apply_cookie_header(req: urllib.request.Request) -> None:
    hv = _cookie_header_value()
    if hv:
        req.add_header("Cookie", hv)


def _urlopen_with_cookies(req: urllib.request.Request, timeout: int):
    opener = _get_cookie_opener()
    if opener is not None:
        return opener.open(req, timeout=timeout)
    _apply_cookie_header(req)
    return urllib.request.urlopen(req, timeout=timeout, context=CTX)


def _curl_cookie_args() -> list[str]:
    fp = _cookie_file_path()
    if fp:
        return ["-b", fp]
    hv = _cookie_header_value()
    if hv:
        return ["-b", hv]
    return []


def find_dois_in_md(text: str) -> list[str]:
    out: list[str] = []
    for m in DOI_RE.finditer(text):
        d = m.group(1).rstrip(").,;")
        if d not in out:
            out.append(d)
    return out


# --- Author–year citations (XLSX ``Author (date)`` column) → DOI → PDF ---

_AUTHOR_YEAR_TAIL = re.compile(r"\((19\d{2}|20\d{2})\)\s*$")
# Trailing "et al." / ", et al" (not used in API searches or as a fake surname).
_ET_AL_TRAIL = re.compile(r"(?i)(?:,)?\s*et\s+al\.?\s*$")
# Co-author separators within one comma-separated chunk (``\b`` avoids splitting ``Sanderson``).
_AUTHOR_CONJUNCTION_RE = re.compile(r"\s*(?:\band\b|&)\s*", re.I)


def _strip_trailing_et_al(s: str) -> str:
    t = (s or "").strip()
    while True:
        n = _ET_AL_TRAIL.sub("", t).strip().rstrip(",").strip()
        if n == t:
            return n
        t = n


def _split_coauthor_segment(segment: str) -> list[str]:
    """Split one comma-separated chunk on ``and`` / ``&`` between surnames."""
    seg = _strip_trailing_et_al((segment or "").strip())
    if not seg:
        return []
    names = [_strip_trailing_et_al(s.strip()) for s in _AUTHOR_CONJUNCTION_RE.split(seg)]
    return [s for s in names if s]


@dataclass
class ParsedCitation:
    surnames: list[str]
    year: int
    raw: str


@dataclass
class ResolutionResult:
    doi: str | None
    title: str | None
    source: str  # crossref | openalex | scholarly | ""
    status: str  # resolved | unresolved | ambiguous | ambiguous (>5)
    detail: str
    authors: str | None = None
    publication_year: int | None = None
    journal: str | None = None
    ambiguous_candidates: tuple[ResolutionResult, ...] | None = None


def parse_author_year_citation(text: str) -> ParsedCitation | None:
    """Parse ``Surname1, Surname2 and Surname3 (2003)`` style strings."""
    t = (text or "").strip()
    if not t:
        return None
    m = _AUTHOR_YEAR_TAIL.search(t)
    if not m:
        return None
    year = int(m.group(1))
    author = _strip_trailing_et_al(t[: m.start()].strip().rstrip(",").strip())
    if not author:
        return None
    parts = [_strip_trailing_et_al(p.strip()) for p in author.split(",") if p.strip()]
    parts = [p for p in parts if p]
    if not parts:
        return None
    surnames: list[str] = []
    for p in parts:
        surnames.extend(_split_coauthor_segment(p))
    junk = frozenset(
        x.lower() for x in ("et", "al", "etal", "et al", "et al.", "etal.")
    )
    surnames = [s for s in surnames if s and s.strip().lower() not in junk]
    if not surnames:
        return None
    return ParsedCitation(surnames=surnames, year=year, raw=t)


def _mailto_ua() -> str:
    return f"{USER_AGENT} mailto:{_unpaywall_email()}"


def _crossref_item_year(item: dict) -> int | None:
    for key in ("published-print", "published-online", "issued"):
        d = item.get(key) or {}
        parts = d.get("date-parts") or []
        if parts and parts[0]:
            try:
                return int(parts[0][0])
            except (TypeError, ValueError):
                continue
    return None


def _crossref_item_title(item: dict) -> str:
    t = item.get("title")
    if isinstance(t, list) and t:
        return str(t[0])
    if isinstance(t, str):
        return t
    return ""


def _crossref_format_authors(item: dict) -> str:
    parts: list[str] = []
    for a in item.get("author") or []:
        if not isinstance(a, dict):
            continue
        fam = a.get("family")
        given = a.get("given")
        name_lit = a.get("name")
        if isinstance(fam, str) and fam.strip():
            if isinstance(given, str) and given.strip():
                parts.append(f"{given.strip()} {fam.strip()}")
            else:
                parts.append(fam.strip())
        elif isinstance(name_lit, str) and name_lit.strip():
            parts.append(name_lit.strip())
    return "; ".join(parts)


def _crossref_item_journal(item: dict) -> str:
    ct = item.get("container-title")
    if isinstance(ct, list) and ct:
        s = str(ct[0]).strip()
        if s:
            return s
    if isinstance(ct, str) and ct.strip():
        return ct.strip()
    st = item.get("short-container-title")
    if isinstance(st, list) and st:
        s = str(st[0]).strip()
        if s:
            return s
    if isinstance(st, str) and st.strip():
        return st.strip()
    pub = item.get("publisher")
    if isinstance(pub, str) and pub.strip():
        return pub.strip()
    return ""


def _crossref_work_bibliographic_meta(item: dict) -> tuple[str, int | None, str]:
    auth = _crossref_format_authors(item)
    j = _crossref_item_journal(item).strip()
    iy = _crossref_item_year(item)
    return (auth, iy, j)


def _crossref_pool_tied_items(
    pool: list[tuple[int, int, int, dict]],
    parsed_year: int,
    top_year: int,
    top_ov: int,
    top_kw: int,
) -> list[dict]:
    """Works tied with ``pool[0]`` on score triple; publication year must match ``parsed_year``."""
    out: list[dict] = []
    for ym, ov, kw, it in pool:
        if (ym, ov, kw) != (top_year, top_ov, top_kw):
            break
        if _crossref_item_year(it) != parsed_year:
            continue
        out.append(it)
    return out


def _dedupe_crossref_items_by_doi(items: list[dict]) -> list[dict]:
    seen: set[str] = set()
    deduped: list[dict] = []
    for it in items:
        d = _normalize_doi(it.get("DOI"))
        if not d or d in seen:
            continue
        seen.add(d)
        deduped.append(it)
    return deduped


def _candidate_result_crossref(
    it: dict,
    *,
    overlap: int,
    kw_hits: int,
    rank: int,
    total: int,
) -> ResolutionResult:
    au, pub_y, journ = _crossref_work_bibliographic_meta(it)
    return ResolutionResult(
        _normalize_doi(it.get("DOI")),
        _crossref_item_title(it) or None,
        "crossref",
        "ambiguous",
        f"tied candidate {rank}/{total}; overlap={overlap} kw_hits={kw_hits}",
        authors=au or None,
        publication_year=pub_y,
        journal=journ or None,
    )


# Crossref/OpenAlex often use Unicode hyphens (e.g. U+2010) in ``family``; citations use ASCII "-".
_AUTHOR_CMP_TRANS = {ord(c): ord("-") for c in "\u2010\u2011\u2012\u2013\u2014\u2212"}


def _normalize_author_cmp(s: str) -> str:
    """Lowercase + map unicode dash/hyphen variants to ASCII for surname equality checks."""
    return (s or "").strip().lower().translate(_AUTHOR_CMP_TRANS)


def _crossref_item_families(item: dict) -> set[str]:
    out: set[str] = set()
    for a in item.get("author") or []:
        fam = a.get("family")
        if fam and isinstance(fam, str):
            out.add(_normalize_author_cmp(fam))
    return out


def _crossref_surname_overlap(parsed: ParsedCitation, item: dict) -> int:
    fams = _crossref_item_families(item)
    return sum(1 for s in parsed.surnames if _normalize_author_cmp(s) in fams)


def _crossref_first_listed_family(item: dict) -> str | None:
    """Family name of Crossref's first-listed author (``sequence: first`` when present)."""
    authors = item.get("author") or []
    if not authors:
        return None
    for a in authors:
        if str(a.get("sequence", "")).lower() == "first":
            fam = a.get("family")
            if isinstance(fam, str) and fam.strip():
                return fam.strip()
    fam0 = authors[0].get("family")
    if isinstance(fam0, str) and fam0.strip():
        return fam0.strip()
    return None


def _crossref_first_author_matches_citation(parsed: ParsedCitation, item: dict) -> bool:
    """Citation's first surname must equal Crossref's first author ``family`` (case-insensitive)."""
    if not parsed.surnames:
        return True
    want = _normalize_author_cmp(parsed.surnames[0])
    cf = _crossref_first_listed_family(item)
    if not cf:
        return False
    return _normalize_author_cmp(cf) == want


# Minimal stopwords for topic tokens taken from optional XLSX context columns.
_BOOST_STOP = frozenset(
    "a an the and or but if in on at to for of as by is was are were be been "
    "being from with into through over under than then this that these those "
    "their our your its his her their them they we you it not no yes all any "
    "each every both few more most some such only own same so than too very "
    "can could should would may might must will shall about also into per "
    "between among within without against during before after above below "
    "across near far".split()
)


def title_boost_tokens(text: str) -> frozenset[str]:
    """Lowercase alpha tokens (length >= 4) from free text, minus common stopwords."""
    raw = re.findall(r"[A-Za-z]{4,}", (text or "").lower())
    return frozenset(t for t in raw if t not in _BOOST_STOP)


def _keyword_hits_in_lower_title(title_lower: str, boost: frozenset[str]) -> int:
    if not boost or not title_lower:
        return 0
    return sum(1 for kw in boost if kw in title_lower)


def _normalize_doi(doi: str | None) -> str | None:
    if not doi or not isinstance(doi, str):
        return None
    d = doi.strip()
    if d.lower().startswith("https://doi.org/"):
        d = d[16:]
    elif d.lower().startswith("http://doi.org/"):
        d = d[15:]
    return d or None


def _doi_from_scan_fragment(fragment: str) -> str | None:
    """Best-effort DOI from regex/HTML substring (e.g. SERP snippets)."""
    m = re.match(
        r"(10\.\d{4,9}/[A-Za-z0-9./\-_:;%]+)",
        (fragment or "").strip(),
        re.I,
    )
    return _normalize_doi(m.group(1)) if m else None


def _env_truthy(name: str) -> bool:
    return os.environ.get(name, "").strip().lower() in ("1", "true", "yes")


def _crossref_resolve_citation_phase(
    parsed: ParsedCitation,
    boost: frozenset[str],
    *,
    journal_only: bool,
) -> ResolutionResult:
    """Single Crossref /works request and scoring (journal-article filter optional)."""
    bib = " ".join(parsed.surnames) + f" {parsed.year}"
    q = urllib.parse.quote(bib, safe="")
    mail = urllib.parse.quote(_unpaywall_email(), safe="")
    y = parsed.year
    filt = f"from-pub-date:{y}-01-01,until-pub-date:{y}-12-31"
    if journal_only:
        filt = filt + ",type:journal-article"
    url = (
        "https://api.crossref.org/works?"
        f"query.bibliographic={q}&rows=20&mailto={mail}"
        f"&filter={urllib.parse.quote(filt, safe=',:-')}"
    )
    if parsed.surnames:
        author_q = urllib.parse.quote(parsed.surnames[0], safe="")
        url += f"&query.author={author_q}"
    try:
        data, _ = http_get_urllib(
            url, extra_headers={"User-Agent": _mailto_ua()}
        )
        msg = json.loads(data.decode("utf-8", errors="replace"))["message"]
    except Exception as e:
        return ResolutionResult(
            None, None, "", "unresolved", f"crossref request failed: {e}"
        )
    items = msg.get("items") or []
    if not items:
        phase = "journal-article" if journal_only else "all types"
        return ResolutionResult(
            None, None, "", "unresolved", f"crossref: no items ({phase})"
        )
    n_raw = len(items)
    items = [it for it in items if _crossref_first_author_matches_citation(parsed, it)]
    dbg = _env_truthy("AI4RESEARCH_CROSSREF_DEBUG")
    if dbg:
        ph = "journal-article" if journal_only else "all-types"
        print(
            f"[crossref-debug] phase={ph} api_items={n_raw} "
            f"after_first_author_filter={len(items)}",
            file=sys.stderr,
        )
        print(f"[crossref-debug] GET {url}", file=sys.stderr)
    if not items:
        fa = parsed.surnames[0] if parsed.surnames else "?"
        phase = "journal-article" if journal_only else "all types"
        return ResolutionResult(
            None,
            None,
            "",
            "unresolved",
            f"crossref: no hit whose first author family matches citation first {fa!r} ({phase})",
        )

    scored: list[tuple[int, int, int, dict]] = []
    for it in items:
        iy = _crossref_item_year(it)
        ov = _crossref_surname_overlap(parsed, it)
        year_match = 1 if iy == parsed.year else 0
        tit = _crossref_item_title(it).lower()
        kw = _keyword_hits_in_lower_title(tit, boost)
        scored.append((year_match, ov, kw, it))
    with_year = [x for x in scored if x[0] == 1]
    pool = with_year if with_year else scored
    pool.sort(key=lambda x: (x[0], x[1], x[2]), reverse=True)
    if dbg:
        ph = "journal-article" if journal_only else "all-types"
        print(
            f"[crossref-debug] phase={ph} ranked rows ({len(pool)}), "
            f"sort=(year_match, overlap, kw_hits) desc:",
            file=sys.stderr,
        )
        for i, (ym, ov, kw, it) in enumerate(pool, start=1):
            d = _normalize_doi(it.get("DOI")) or "?"
            tit = (_crossref_item_title(it) or "")[:80]
            iy = _crossref_item_year(it)
            print(
                f"  {i:2} year_match={ym} overlap={ov} kw_hits={kw} "
                f"crossref_year={iy} doi={d}",
                file=sys.stderr,
            )
            print(f"      title={tit!r}", file=sys.stderr)
    if not pool:
        return ResolutionResult(
            None, None, "", "unresolved", "crossref: empty pool"
        )
    top_year, top_ov, top_kw, top_it = pool[0]
    au, pub_y, journ = _crossref_work_bibliographic_meta(top_it)
    if top_ov < 1:
        return ResolutionResult(
            None,
            _crossref_item_title(top_it) or None,
            "",
            "unresolved",
            f"crossref: weak author overlap ({top_ov})",
            authors=au or None,
            publication_year=pub_y,
            journal=journ or None,
        )
    doi = _normalize_doi(top_it.get("DOI"))
    title = _crossref_item_title(top_it) or None
    tied_raw = _crossref_pool_tied_items(
        pool, parsed.year, top_year, top_ov, top_kw
    )
    tied_dedup = _dedupe_crossref_items_by_doi(tied_raw)
    if dbg and len(tied_dedup) >= 2:
        ph = "journal-article" if journal_only else "all-types"
        dois_all = [
            _normalize_doi(x.get("DOI")) for x in tied_dedup if _normalize_doi(x.get("DOI"))
        ]
        print(
            f"[crossref-debug] phase={ph} tied_at_top overlap={top_ov} kw={top_kw} "
            f"n_distinct_doi={len(tied_dedup)}: {dois_all}",
            file=sys.stderr,
        )
    if len(tied_dedup) >= 2:
        n_sur = len(parsed.surnames)
        if n_sur >= 2 and top_ov == 1:
            return ResolutionResult(
                None,
                None,
                "",
                "unresolved",
                f"crossref: {len(tied_dedup)} hits tied at overlap=1; citation lists "
                f"{n_sur} surnames but results only match '{parsed.surnames[0]}' "
                f"(likely unrelated works e.g. columns). Try --scholarly or add "
                f"distinct title words via context columns / manual DOI.",
            )
        if len(tied_dedup) > 5:
            dois_sample = ", ".join(
                repr(_normalize_doi(x.get("DOI"))) for x in tied_dedup[:3]
            )
            return ResolutionResult(
                None,
                None,
                "",
                "ambiguous (>5)",
                f"crossref: {len(tied_dedup)} works tied overlap={top_ov} kw={top_kw} "
                f"(sample DOIs {dois_sample}…)",
            )
        cands = tuple(
            _candidate_result_crossref(
                it,
                overlap=top_ov,
                kw_hits=top_kw,
                rank=i,
                total=len(tied_dedup),
            )
            for i, it in enumerate(tied_dedup, start=1)
        )
        return ResolutionResult(
            None,
            None,
            "crossref",
            "ambiguous",
            f"crossref: {len(tied_dedup)} tied matches overlap={top_ov} kw={top_kw}",
            ambiguous_candidates=cands,
        )
    if not doi:
        return ResolutionResult(
            None,
            title,
            "",
            "unresolved",
            "crossref: top item missing DOI",
            authors=au or None,
            publication_year=pub_y,
            journal=journ or None,
        )
    kw_note = f" kw_hits={top_kw}" if boost else ""
    first_note = (
        f" first_author={parsed.surnames[0]!r}" if parsed.surnames else ""
    )
    return ResolutionResult(
        doi,
        title,
        "crossref",
        "resolved",
        f"overlap={top_ov} year_match={top_year}{kw_note}{first_note}",
        authors=au or None,
        publication_year=pub_y,
        journal=journ or None,
    )


def crossref_resolve_citation(
    parsed: ParsedCitation,
    boost_tokens: frozenset[str] | None = None,
) -> ResolutionResult:
    """Query Crossref works API; pick best item by year + author overlap + optional title keywords.

    Only works whose **first-listed Crossref author** ``family`` equals the **first surname**
    in the citation cell (case-insensitive) are considered; others are dropped. The request
    adds ``query.author`` for that first surname when present, and ``filter`` with
    ``from-pub-date`` / ``until-pub-date`` for the citation year (Crossref has no
    ``publication_year`` key; this scopes results to that calendar year).

    Search runs **twice**: first with ``type:journal-article``; if that pass does not
    resolve and does not yield an ambiguous candidate list, the same query is repeated
    without that type filter so book chapters and other DOI types can match.
    """
    boost = boost_tokens or frozenset()
    # Query authors + year only. Theme words from --keyword-boost / context columns are
    # *not* appended here: they skew Crossref's ranked list and can push the right paper
    # (e.g. Milligan + Stabile + year) out of the first page of results.
    journal = _crossref_resolve_citation_phase(parsed, boost, journal_only=True)
    if journal.status == "resolved" or journal.ambiguous_candidates:
        return journal
    broad = _crossref_resolve_citation_phase(parsed, boost, journal_only=False)
    if broad.status in ("resolved", "ambiguous", "ambiguous (>5)") and broad.detail:
        suffix = "; crossref_types=expanded"
        if suffix not in broad.detail:
            broad = ResolutionResult(
                broad.doi,
                broad.title,
                broad.source,
                broad.status,
                broad.detail + suffix,
                authors=broad.authors,
                publication_year=broad.publication_year,
                journal=broad.journal,
                ambiguous_candidates=broad.ambiguous_candidates,
            )
    return broad


def _openalex_format_authors(work: dict) -> str:
    parts: list[str] = []
    for a in work.get("authorships") or []:
        if not isinstance(a, dict):
            continue
        auth = a.get("author") or {}
        disp = (auth.get("display_name") or "") or (
            auth.get("author_display_name") or ""
        )
        if isinstance(disp, str) and disp.strip():
            parts.append(disp.strip())
    return "; ".join(parts)


def _openalex_item_journal(work: dict) -> str:
    def from_src(src: object) -> str:
        if isinstance(src, dict):
            n = src.get("display_name")
            if isinstance(n, str) and n.strip():
                return n.strip()
        return ""

    pl = work.get("primary_location")
    if isinstance(pl, dict):
        j = from_src(pl.get("source"))
        if j:
            return j
    for loc in work.get("locations") or []:
        if isinstance(loc, dict):
            j = from_src(loc.get("source"))
            if j:
                return j
    bib = work.get("biblio")
    if isinstance(bib, dict):
        jn = bib.get("journal") or bib.get("repository")
        if isinstance(jn, str) and jn.strip():
            return jn.strip()
    return ""


def _openalex_work_bibliographic_meta(work: dict) -> tuple[str, int | None, str]:
    au = _openalex_format_authors(work)
    py_raw = work.get("publication_year")
    pub_y: int | None
    try:
        pub_y = int(py_raw) if py_raw is not None else None
    except (TypeError, ValueError):
        pub_y = None
    j = _openalex_item_journal(work).strip()
    return (au, pub_y, j)


def _openalex_scored_tied_works(
    scored: list[tuple[int, int, dict]],
    parsed_year: int,
    top_o: int,
    top_kw: int,
) -> list[dict]:
    """Works tied with ``scored[0]`` on (surname overlap, kw_hits); pub year must match."""
    out: list[dict] = []
    for o, kw, w in scored:
        if (o, kw) != (top_o, top_kw):
            break
        py_raw = w.get("publication_year")
        try:
            if int(py_raw) != parsed_year:
                continue
        except (TypeError, ValueError):
            continue
        out.append(w)
    return out


def _dedupe_openalex_works_by_doi(works: list[dict]) -> list[dict]:
    seen: set[str] = set()
    deduped: list[dict] = []
    for w in works:
        d = _normalize_doi(w.get("doi"))
        if not d or d in seen:
            continue
        seen.add(d)
        deduped.append(w)
    return deduped


def _candidate_result_openalex(
    work: dict,
    *,
    overlap: int,
    kw_hits: int,
    rank: int,
    total: int,
) -> ResolutionResult:
    au, pub_y, journ = _openalex_work_bibliographic_meta(work)
    return ResolutionResult(
        _normalize_doi(work.get("doi")),
        work.get("display_name") or None,
        "openalex",
        "ambiguous",
        f"tied candidate {rank}/{total}; overlap={overlap} kw_hits={kw_hits}",
        authors=au or None,
        publication_year=pub_y,
        journal=journ or None,
    )


def openalex_resolve_citation(
    parsed: ParsedCitation,
    boost_tokens: frozenset[str] | None = None,
) -> ResolutionResult:
    """Fallback: OpenAlex work search by year + surnames; topic tokens only for title scoring."""
    boost = boost_tokens or frozenset()
    search = " ".join(parsed.surnames)
    if not search.strip():
        return ResolutionResult(None, None, "", "unresolved", "openalex: empty search")
    q = urllib.parse.quote(search, safe="")
    url = (
        "https://api.openalex.org/works?"
        f"filter=publication_year:{parsed.year}&search={q}&per_page=10"
    )
    try:
        data, _ = http_get_urllib(url, openalex_mailto=True)
        msg = json.loads(data.decode("utf-8", errors="replace"))
    except Exception as e:
        return ResolutionResult(
            None, None, "", "unresolved", f"openalex request failed: {e}"
        )
    results = msg.get("results") or []
    if not results:
        return ResolutionResult(None, None, "", "unresolved", "openalex: no results")

    def surname_hits(w: dict) -> int:
        parts: list[str] = []
        for a in w.get("authorships") or []:
            auth = a.get("author") or {}
            disp = (auth.get("display_name") or "") or (
                a.get("author_display_name") or ""
            )
            if disp:
                parts.append(disp.lower())
        blob = " ".join(parts)
        blob_n = _normalize_author_cmp(blob)
        return sum(1 for s in parsed.surnames if _normalize_author_cmp(s) in blob_n)

    def kw_hits(w: dict) -> int:
        disp = (w.get("display_name") or "").lower()
        return _keyword_hits_in_lower_title(disp, boost)

    scored = [(surname_hits(w), kw_hits(w), w) for w in results]
    scored.sort(key=lambda x: (x[0], x[1]), reverse=True)
    top_o, top_kw, top_w = scored[0]
    oa_au, oa_y, oa_j = _openalex_work_bibliographic_meta(top_w)
    if top_o < 1:
        return ResolutionResult(
            None,
            (top_w.get("display_name") or None),
            "",
            "unresolved",
            f"openalex: weak overlap ({top_o})",
            authors=oa_au or None,
            publication_year=oa_y,
            journal=oa_j or None,
        )
    if len(parsed.surnames) >= 2 and top_o < 2:
        return ResolutionResult(
            None,
            (top_w.get("display_name") or None),
            "",
            "unresolved",
            f"openalex: overlap={top_o} for {len(parsed.surnames)}-author citation "
            f"(need ≥2 matching surnames on record)",
            authors=oa_au or None,
            publication_year=oa_y,
            journal=oa_j or None,
        )
    doi = _normalize_doi(top_w.get("doi"))
    title = top_w.get("display_name") or None
    tied_raw = _openalex_scored_tied_works(
        scored, parsed.year, top_o, top_kw
    )
    tied_dedup = _dedupe_openalex_works_by_doi(tied_raw)
    if len(tied_dedup) >= 2:
        n_sur = len(parsed.surnames)
        if n_sur >= 2 and top_o == 1:
            return ResolutionResult(
                None,
                None,
                "",
                "unresolved",
                f"openalex: {len(tied_dedup)} hits tied at overlap=1; citation lists "
                f"{n_sur} surnames",
            )
        if len(tied_dedup) > 5:
            dois_sample = ", ".join(
                repr(_normalize_doi(w.get("doi"))) for w in tied_dedup[:3]
            )
            return ResolutionResult(
                None,
                None,
                "",
                "ambiguous (>5)",
                f"openalex: {len(tied_dedup)} works tied overlap={top_o} kw={top_kw} "
                f"(sample DOIs {dois_sample}…)",
            )
        cands = tuple(
            _candidate_result_openalex(
                w,
                overlap=top_o,
                kw_hits=top_kw,
                rank=i,
                total=len(tied_dedup),
            )
            for i, w in enumerate(tied_dedup, start=1)
        )
        return ResolutionResult(
            None,
            None,
            "openalex",
            "ambiguous",
            f"openalex: {len(tied_dedup)} tied matches overlap={top_o} kw={top_kw}",
            ambiguous_candidates=cands,
        )
    if not doi:
        return ResolutionResult(
            None,
            title,
            "",
            "unresolved",
            "openalex: top result missing doi",
            authors=oa_au or None,
            publication_year=oa_y,
            journal=oa_j or None,
        )
    kw_note = f" kw_hits={top_kw}" if boost else ""
    return ResolutionResult(
        doi,
        title,
        "openalex",
        "resolved",
        f"overlap={top_o}{kw_note}",
        authors=oa_au or None,
        publication_year=oa_y,
        journal=oa_j or None,
    )


_RAW_DOI_IN_TEXT = re.compile(
    r"\b(10\.\d{4,9}/[^\s\],;\)\]\"'>]+)", re.I
)


def _extract_first_doi_from_obj(obj: object, *, max_depth: int = 12) -> str | None:
    """Walk dict/list/str and return first DOI-like substring (normalized)."""
    if max_depth < 0:
        return None
    if isinstance(obj, str):
        m = _RAW_DOI_IN_TEXT.search(obj.replace("%2F", "/"))
        return _normalize_doi(m.group(1)) if m else None
    if isinstance(obj, dict):
        for v in obj.values():
            d = _extract_first_doi_from_obj(v, max_depth=max_depth - 1)
            if d:
                return d
    if isinstance(obj, (list, tuple)):
        for v in obj:
            d = _extract_first_doi_from_obj(v, max_depth=max_depth - 1)
            if d:
                return d
    return None


def _scholarly_bib_title(pub: dict) -> str | None:
    bib = pub.get("bib") or {}
    t = bib.get("title")
    return str(t) if t else None


def _scholarly_pub_bibliographic_meta(pub: dict) -> tuple[str, int | None, str]:
    bib = pub.get("bib") or {}
    auth = bib.get("author")
    if isinstance(auth, str) and auth.strip():
        authors = auth.strip()
    elif isinstance(auth, list):
        authors = "; ".join(str(x).strip() for x in auth if str(x).strip())
    else:
        authors = ""
    py = bib.get("pub_year")
    pub_y: int | None
    try:
        pub_y = int(py) if py is not None else None
    except (TypeError, ValueError):
        pub_y = None
    venue = bib.get("venue") or bib.get("journal") or ""
    journal = venue.strip() if isinstance(venue, str) else ""
    return (authors, pub_y, journal)


def _title_contains_all_tokens(title: str | None, tokens: frozenset[str]) -> bool:
    """Every token must appear as a substring in ``title`` (case-insensitive)."""
    if not tokens:
        return True
    t = (title or "").lower()
    return all(kw in t for kw in tokens)


def scholarly_resolve_citation(
    parsed: ParsedCitation,
    boost_tokens: frozenset[str] | None = None,
    *,
    keyword_boost_required: frozenset[str] | None = None,
) -> ResolutionResult:
    """
    Query Google Scholar via optional `scholarly` package after Crossref did not resolve.

    The search string includes **every** token in ``boost_tokens`` at once (space-separated)
    so Scholar narrows on all topic words together. If ``keyword_boost_required`` is
    non-empty (from ``--keyword-boost``), only hits whose **title** contains **all**
    of those terms are considered for DOI extraction.

    Returns a DOI only when one can be extracted from Scholar metadata (no PDF scraping).
    """
    try:
        from scholarly import scholarly as scholarly_api
    except ImportError:
        return ResolutionResult(
            None,
            None,
            "",
            "unresolved",
            "scholarly: not installed (pip install scholarly)",
        )

    boost = boost_tokens or frozenset()
    required_in_title = keyword_boost_required or frozenset()
    q = " ".join(parsed.surnames) + f" {parsed.year}"
    if boost:
        # Use every keyword together (no truncation) so Scholar narrows the query.
        q = f"{q} {' '.join(sorted(boost))}".strip()
    try:
        it = scholarly_api.search_pubs(q)
        pubs: list[dict] = []
        for _i, p in enumerate(it):
            if isinstance(p, dict):
                pubs.append(p)
            else:
                try:
                    pubs.append(dict(p))  # type: ignore[arg-type]
                except (TypeError, ValueError):
                    continue
            if len(pubs) >= 8:
                break
    except Exception as e:
        time.sleep(1.5)
        return ResolutionResult(
            None, None, "", "unresolved", f"scholarly: search failed: {e}"
        )
    if not pubs:
        time.sleep(1.5)
        return ResolutionResult(None, None, "", "unresolved", "scholarly: no results")

    def score_pub(pub: dict) -> int:
        bib = pub.get("bib") or {}
        sc = 0
        py = bib.get("pub_year")
        try:
            if py is not None and int(py) == parsed.year:
                sc += 10
        except (TypeError, ValueError):
            pass
        auth = bib.get("author")
        blob = ""
        if isinstance(auth, str):
            blob = auth.lower()
        elif isinstance(auth, list):
            blob = " ".join(str(x) for x in auth).lower()
        for s in parsed.surnames:
            if s.lower() in blob:
                sc += 3
        tit = (bib.get("title") or "").lower()
        for kw in boost:
            if len(kw) >= 4 and kw in tit:
                sc += 2
        return sc

    ranked = sorted(pubs, key=score_pub, reverse=True)
    if required_in_title:
        ranked = [
            p
            for p in ranked
            if _title_contains_all_tokens(_scholarly_bib_title(p), required_in_title)
        ]
        if not ranked:
            time.sleep(2.0)
            return ResolutionResult(
                None,
                None,
                "",
                "unresolved",
                "scholarly: no result title contains all --keyword-boost terms",
            )
    for pub in ranked[:3]:
        doi = _extract_first_doi_from_obj(pub)
        tit = _scholarly_bib_title(pub)
        if doi:
            sa, sy, sj = _scholarly_pub_bibliographic_meta(pub)
            time.sleep(2.0)
            return ResolutionResult(
                doi,
                tit,
                "scholarly",
                "resolved",
                "Google Scholar preview metadata",
                authors=sa or None,
                publication_year=sy,
                journal=sj or None,
            )

    best = ranked[0]
    try:
        filled = scholarly_api.fill(best)
        if not isinstance(filled, dict):
            try:
                filled = dict(filled)  # type: ignore[arg-type]
            except (TypeError, ValueError):
                filled = best
    except Exception as e:
        time.sleep(2.0)
        ba, by_, bj = _scholarly_pub_bibliographic_meta(best)
        return ResolutionResult(
            None,
            _scholarly_bib_title(best),
            "",
            "unresolved",
            f"scholarly: no DOI in preview; fill failed: {e}",
            authors=ba or None,
            publication_year=by_,
            journal=bj or None,
        )
    doi = _extract_first_doi_from_obj(filled)
    tit = _scholarly_bib_title(filled)
    time.sleep(2.0)
    if doi:
        fa, fy, fj = _scholarly_pub_bibliographic_meta(filled)
        return ResolutionResult(
            doi,
            tit,
            "scholarly",
            "resolved",
            "Google Scholar filled metadata",
            authors=fa or None,
            publication_year=fy,
            journal=fj or None,
        )
    fa, fy, fj = _scholarly_pub_bibliographic_meta(filled)
    return ResolutionResult(
        None,
        tit,
        "",
        "unresolved",
        "scholarly: no DOI in top Scholar hits",
        authors=fa or None,
        publication_year=fy,
        journal=fj or None,
    )


def citation_web_search_query_string(
    parsed: ParsedCitation,
    boost_tokens: frozenset[str] | None,
) -> str:
    parts = list(parsed.surnames) + [str(parsed.year)]
    b = boost_tokens or frozenset()
    if b:
        parts.extend(sorted(b))
    return " ".join(parts)


def print_citation_web_search_urls(
    parsed: ParsedCitation,
    boost_tokens: frozenset[str] | None,
    *,
    cite_label: str = "",
) -> None:
    """Print Google / Scholar / DuckDuckGo links (stderr) using the same query shape as Scholarly."""
    q = citation_web_search_query_string(parsed, boost_tokens)
    enc = urllib.parse.quote_plus(q)
    hdr = "[search-urls]"
    if cite_label:
        hdr += f" {cite_label}"
    print(hdr, file=sys.stderr)
    print(f"  Query: {q!r}", file=sys.stderr)
    print(f"  Google:       https://www.google.com/search?q={enc}", file=sys.stderr)
    print(f"  Scholar:      https://scholar.google.com/scholar?q={enc}", file=sys.stderr)
    print(f"  DuckDuckGo:   https://duckduckgo.com/?q={enc}", file=sys.stderr)


def _crossref_work_message_for_doi(doi: str) -> dict | None:
    try:
        url = (
            "https://api.crossref.org/works/"
            + urllib.parse.quote(doi.strip(), safe="")
        )
        data, _ = http_get_urllib(
            url, extra_headers={"User-Agent": _mailto_ua()}
        )
        return json.loads(data.decode("utf-8", errors="replace")).get("message")
    except Exception:
        return None


def try_duckduckgo_lite_resolve(
    parsed: ParsedCitation,
    boost_tokens: frozenset[str] | None,
) -> ResolutionResult | None:
    """Fetch DuckDuckGo Lite HTML; first DOI whose Crossref record passes first-author check wins."""
    q = citation_web_search_query_string(parsed, boost_tokens)
    lite_url = "https://lite.duckduckgo.com/lite/?q=" + urllib.parse.quote_plus(q)
    try:
        body, _ = http_get_urllib(
            lite_url,
            extra_headers={
                "User-Agent": USER_AGENT,
                "Referer": "https://duckduckgo.com/",
            },
        )
    except Exception as e:
        print(f"[duckduckgo-lite] {e}", file=sys.stderr)
        return None
    text = body.decode("utf-8", errors="replace")
    seen: set[str] = set()
    for m in _RAW_DOI_IN_TEXT.finditer(text):
        d = _doi_from_scan_fragment(m.group(1))
        if not d or d in seen:
            continue
        seen.add(d)
        item = _crossref_work_message_for_doi(d)
        if item and _crossref_first_author_matches_citation(parsed, item):
            title = _crossref_item_title(item) or None
            au, pub_y, journ = _crossref_work_bibliographic_meta(item)
            return ResolutionResult(
                d,
                title,
                "duckduckgo-lite",
                "resolved",
                "DOI from DuckDuckGo Lite + Crossref author check",
                authors=au or None,
                publication_year=pub_y,
                journal=journ or None,
            )
    time.sleep(0.25)
    return None


def resolve_citation_to_work(
    parsed: ParsedCitation,
    boost_tokens: frozenset[str] | None = None,
    *,
    use_scholarly: bool = False,
    keyword_boost_required: frozenset[str] | None = None,
    try_web_html_resolve: bool = False,
) -> ResolutionResult:
    cr = crossref_resolve_citation(parsed, boost_tokens)
    if cr.status == "resolved":
        return cr
    if cr.status in ("ambiguous", "ambiguous (>5)"):
        return cr
    if use_scholarly:
        gs = scholarly_resolve_citation(
            parsed,
            boost_tokens,
            keyword_boost_required=keyword_boost_required,
        )
        if gs.status == "resolved":
            return gs
        if gs.status == "ambiguous":
            return gs
    else:
        gs = ResolutionResult(None, None, "", "unresolved", "")
    oa = openalex_resolve_citation(parsed, boost_tokens)
    if oa.status == "resolved":
        return oa
    if oa.status in ("ambiguous", "ambiguous (>5)"):
        return oa
    if try_web_html_resolve:
        ddr = try_duckduckgo_lite_resolve(parsed, boost_tokens)
        if ddr is not None:
            return ddr
    parts = [cr.detail]
    if use_scholarly and gs.detail:
        parts.append(gs.detail)
    parts.append(oa.detail)
    merged = "; then ".join(p for p in parts if p)
    if cr.title or merged:
        return ResolutionResult(
            None,
            cr.title,
            "",
            "unresolved",
            merged,
            authors=cr.authors,
            publication_year=cr.publication_year,
            journal=cr.journal,
        )
    return oa


def _md_escape_cell(s: str) -> str:
    t = (s or "").replace("\n", " ").replace("\r", " ")
    t = " ".join(t.split())
    return t.replace("|", r"\|")


def _markdown_table_row(cells: list[str]) -> str:
    return "| " + " | ".join(_md_escape_cell(c) for c in cells) + " |"


_XLSX_DOWNLOAD_SUMMARY_COLUMNS = (
    "citation",
    "authors",
    "publication_year",
    "journal",
    "title",
    "doi",
    "status",
    "detail",
    "source",
)


def write_xlsx_download_csv(
    *,
    csv_path: Path,
    sheet_tables: list[tuple[str, list[list[str]]]],
) -> None:
    """One UTF-8 CSV: ``sheet`` column plus the same columns as the markdown tables."""
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    cols = list(_XLSX_DOWNLOAD_SUMMARY_COLUMNS)
    ncols = len(cols)
    header = ["sheet"] + cols
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        w.writerow(header)
        for sheet_name, rows in sheet_tables:
            if not rows:
                continue
            for r in rows:
                row = list(r)
                while len(row) < ncols:
                    row.append("")
                w.writerow([sheet_name] + row[:ncols])


def build_xlsx_download_markdown(
    *,
    xlsx_path: Path,
    sheet_tables: list[tuple[str, list[list[str]]]],
    intro_extra: str = "",
) -> str:
    """``sheet_tables`` = (sheet_name, rows of 9 columns) excluding header row."""
    lines: list[str] = [
        "# XLSX citation download summary",
        "",
        f"Source: `{xlsx_path}`",
        intro_extra.rstrip(),
        "",
    ]
    header = list(_XLSX_DOWNLOAD_SUMMARY_COLUMNS)
    ncols = len(header)
    for sheet_name, rows in sheet_tables:
        if not rows:
            continue
        lines.append(f"## Sheet: {sheet_name}")
        lines.append("")
        lines.append(_markdown_table_row(header))
        lines.append("| " + " | ".join(["---"] * ncols) + " |")
        for r in rows:
            row = list(r)
            while len(row) < ncols:
                row.append("")
            lines.append(_markdown_table_row(row[:ncols]))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def iter_xlsx_author_date_citations(
    xlsx_path: Path,
    citation_column: str,
    context_column_names: tuple[str, ...] = (),
) -> tuple[list[tuple[str, int, str, str]], list[str]]:
    """
    Return (list of (sheet_name, excel_row_number, citation_text, context_text), warnings).

    ``context_text`` is a single string built from same-row cells whose headers
    (exact strip match) appear in ``context_column_names``. Row numbers are 1-based.
    """
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError:
        raise SystemExit(
            "XLSX mode requires openpyxl. Install: pip install openpyxl"
        ) from None
    want = citation_column.strip()
    ctx_wants = tuple(c.strip() for c in context_column_names if c and c.strip())
    out: list[tuple[str, int, str, str]] = []
    warns: list[str] = []
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        for sheet in wb.worksheets:
            col_idx: int | None = None
            ctx_indices: dict[str, int] = {}
            data_rows: list[tuple[int, str, str]] = []
            state = "seek_header"
            for r_i, row in enumerate(
                sheet.iter_rows(values_only=True), start=1
            ):
                if state == "seek_header":
                    if not row:
                        continue
                    hdr: dict[str, int] = {}
                    for c_i, val in enumerate(row):
                        if val is None:
                            continue
                        s = str(val).strip()
                        if s:
                            hdr[s] = c_i
                    if want not in hdr:
                        continue
                    col_idx = hdr[want]
                    for cname in ctx_wants:
                        if cname in hdr:
                            ctx_indices[cname] = hdr[cname]
                        else:
                            warns.append(
                                f"sheet {sheet.title!r}: context column header "
                                f"{cname!r} not found (same row as {want!r})"
                            )
                    state = "data"
                    continue
                if state == "data":
                    if col_idx is None:
                        break
                    vals = list(row)
                    while len(vals) <= col_idx:
                        vals.append(None)
                    cell = vals[col_idx]
                    if cell is None:
                        continue
                    st = str(cell).strip()
                    if not st:
                        continue
                    ctx_parts: list[str] = []
                    for _cname, ci in sorted(
                        ctx_indices.items(), key=lambda x: x[1]
                    ):
                        while len(vals) <= ci:
                            vals.append(None)
                        cv = vals[ci]
                        if cv is not None:
                            cs = str(cv).strip()
                            if cs:
                                ctx_parts.append(cs)
                    ctx_blob = " ".join(ctx_parts)
                    data_rows.append((r_i, st, ctx_blob))
            if col_idx is None:
                warns.append(
                    f"sheet {sheet.title!r}: no column header matching {want!r}"
                )
                continue
            for r_i, st, ctx_blob in data_rows:
                out.append((sheet.title or "Sheet", r_i, st, ctx_blob))
    finally:
        wb.close()
    return out, warns


def slug_from_doi(doi: str) -> str:
    return doi.replace("/", "_").replace("(", "").replace(")", "")


def referer_for_url(url: str) -> str:
    p = urllib.parse.urlparse(url)
    if p.scheme and p.netloc:
        return f"{p.scheme}://{p.netloc}/"
    return "https://www.google.com/"


def _http_get_urllib_once(
    url: str,
    *,
    user_agent: str | None = None,
    extra_headers: dict[str, str] | None = None,
) -> tuple[bytes, str]:
    h = {
        "User-Agent": user_agent or USER_AGENT,
        "Accept": "text/html,application/xhtml+xml,application/pdf;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
        "Accept-Encoding": "gzip, deflate",
        "Sec-Fetch-Site": "none",
        "Sec-Fetch-Mode": "navigate",
        "Sec-Fetch-Dest": "document",
        "Sec-Fetch-User": "?1",
        "Sec-Ch-Ua": '"Chromium";v="136", "Google Chrome";v="136", "Not.A/Brand";v="99"',
        "Sec-Ch-Ua-Mobile": "?0",
        "Sec-Ch-Ua-Platform": '"Windows"',
    }
    if extra_headers:
        h.update(extra_headers)
    req = urllib.request.Request(url, headers=h, method="GET")
    with _urlopen_with_cookies(req, 90) as resp:
        data = _decompress_if_needed(resp.read(), resp)
        return data, resp.geturl()


def http_get_urllib(
    url: str,
    *,
    openalex_mailto: bool = False,
    extra_headers: dict[str, str] | None = None,
) -> tuple[bytes, str]:
    ua = openalex_user_agent() if openalex_mailto or "api.openalex.org" in url else USER_AGENT
    last_err: Exception | None = None
    for attempt in range(_MAX_RETRIES):
        try:
            return _http_get_urllib_once(
                url,
                user_agent=ua,
                extra_headers=extra_headers,
            )
        except urllib.error.HTTPError as e:
            last_err = e
            if e.code in _RETRYABLE_STATUS and attempt < _MAX_RETRIES - 1:
                time.sleep((2**attempt) * 0.4 + random.random() * 0.3)
                continue
            raise
        except (urllib.error.URLError, TimeoutError, OSError) as e:
            last_err = e
            if attempt < _MAX_RETRIES - 1:
                time.sleep((2**attempt) * 0.3 + random.random() * 0.2)
                continue
            raise
    assert last_err is not None
    raise last_err


def curl_fetch(url: str, referer: str | None = None) -> bytes | None:
    if not CURL:
        return None
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".bin")
    tmp.close()
    path = Path(tmp.name)
    ref = referer or referer_for_url(url)
    try:
        cmd = [
            CURL,
            "-L",
            "-s",
            "--compressed",
            "--max-time",
            "120",
            "-A",
            USER_AGENT,
            "-e",
            ref,
            *_curl_cookie_args(),
            "-o",
            str(path),
            url,
        ]
        # Try with --http2 first, fallback if not supported
        try:
            res = subprocess.run([CURL, "--http2", "--help"], capture_output=True, text=True)
            if res.returncode == 0 and "does not support this" not in res.stderr:
                cmd.insert(3, "--http2")
        except Exception:
            pass

        subprocess.run(
            cmd,
            check=True,
            timeout=130,
        )
        data = path.read_bytes()
        return data if data else None
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
        return None
    finally:
        path.unlink(missing_ok=True)


def http_get(url: str) -> tuple[bytes, str]:
    try:
        return http_get_urllib(url)
    except Exception:
        c = curl_fetch(url)
        if c is None:
            raise
        return c, url


def unpaywall_lookup(doi: str) -> dict | None:
    ck = ("unpaywall", doi.lower().strip())
    if ck in _lookup_json_cache:
        return _lookup_json_cache[ck]
    q = urllib.parse.quote(doi, safe="")
    url = f"https://api.unpaywall.org/v2/{q}?email={urllib.parse.quote(_unpaywall_email())}"
    try:
        _rate_limit_source("unpaywall")
        data, _ = http_get_urllib(url)
        out = json.loads(data.decode("utf-8", errors="replace"))
        _lookup_json_cache[ck] = out
        return out
    except urllib.error.HTTPError as e:
        body = b""
        try:
            body = e.read() or b""
        except Exception:
            pass
        print(
            f"Unpaywall HTTP {e.code} for DOI {doi!r}: {_body_snippet(body)}",
            file=sys.stderr,
        )
        _lookup_json_cache[ck] = None
        return None
    except (urllib.error.URLError, json.JSONDecodeError, TimeoutError, OSError) as e:
        print(f"Unpaywall request failed for DOI {doi!r}: {e}", file=sys.stderr)
        _lookup_json_cache[ck] = None
        return None


def iter_unpaywall_pdf_urls(uw: dict) -> list[str]:
    seen: list[str] = []
    for loc in [uw.get("best_oa_location"), *list(uw.get("oa_locations") or [])]:
        if not loc:
            continue
        for key in ("url_for_pdf", "url"):
            u = loc.get(key)
            if not u:
                continue
            ul = u.lower()
            if ul.endswith(".pdf") or "/pdf" in ul or "pdf?" in ul:
                if u not in seen:
                    seen.append(u)
            elif key == "url_for_pdf":
                if u not in seen:
                    seen.append(u)
    return seen


def unpaywall_landing_follow_urls(uw: dict, direct: list[str]) -> list[str]:
    """URLs that may redirect to PDF (same OA entry `url` when not heuristic-matched)."""
    out: list[str] = []
    seen = set(direct)
    for loc in [uw.get("best_oa_location"), *list(uw.get("oa_locations") or [])]:
        if not loc:
            continue
        u = loc.get("url")
        if not u or u in seen:
            continue
        ul = u.lower()
        if ul.endswith(".pdf") or "/pdf" in ul or "pdf?" in ul:
            continue
        seen.add(u)
        out.append(u)
    return out


def semantic_scholar_lookup(doi: str) -> dict | None:
    global _semantic_scholar_backoff_until
    ck = ("semanticscholar", doi.lower().strip())
    if ck in _lookup_json_cache:
        return _lookup_json_cache[ck]
    if _semantic_scholar_backoff_until > time.time():
        return None
    q = urllib.parse.quote(doi, safe="")
    url = f"https://api.semanticscholar.org/graph/v1/paper/DOI:{q}?fields=openAccessPdf,externalIds"
    try:
        _rate_limit_source("semanticscholar")
        data, _ = http_get_urllib(url)
        out = json.loads(data.decode("utf-8", errors="replace"))
        _lookup_json_cache[ck] = out
        return out
    except urllib.error.HTTPError as e:
        body = b""
        try:
            body = e.read() or b""
        except Exception:
            pass
        print(
            f"Semantic Scholar HTTP {e.code} for DOI {doi!r}: {_body_snippet(body)}",
            file=sys.stderr,
        )
        if e.code == 429:
            # Avoid repeated rate-limited requests for every DOI in one run.
            _semantic_scholar_backoff_until = time.time() + 1800
        _lookup_json_cache[ck] = None
        return None
    except (urllib.error.URLError, json.JSONDecodeError, TimeoutError, OSError) as e:
        print(f"Semantic Scholar request failed for DOI {doi!r}: {e}", file=sys.stderr)
        _lookup_json_cache[ck] = None
        return None


def iter_semantic_scholar_pdf_urls(ss: dict) -> list[str]:
    out: list[str] = []
    oa = ss.get("openAccessPdf")
    if oa and oa.get("url"):
        out.append(oa["url"])
    return out


def preprint_pdf_urls(ss: dict) -> list[str]:
    """Construct PDF URLs from arXiv and PubMed Central IDs found in Semantic Scholar metadata."""
    ids = ss.get("externalIds") or {}
    out: list[str] = []
    arxiv_id = ids.get("ArXiv")
    if arxiv_id:
        out.append(f"https://arxiv.org/pdf/{arxiv_id}.pdf")
    pmc_id = ids.get("PubMedCentral")
    if pmc_id:
        pid = str(pmc_id).strip().replace("PMC", "")
        ep = europepmc_pdf_urls_for_pmc_id(pid)
        out.extend(ep)
        if not ep:
            out.append(f"https://www.ncbi.nlm.nih.gov/pmc/articles/PMC{pid}/pdf/")
    return out


def openalex_pdf_urls(doi: str) -> list[str]:
    ck = ("openalex", doi.lower().strip())
    if ck in _lookup_urls_cache:
        return list(_lookup_urls_cache[ck])
    q = urllib.parse.quote(doi, safe="")
    url = f"https://api.openalex.org/works/https://doi.org/{q}"
    try:
        _rate_limit_source("openalex")
        data, _ = http_get_urllib(url)
        msg = json.loads(data.decode("utf-8", errors="replace"))
    except Exception:
        _lookup_urls_cache[ck] = ()
        return []
    out: list[str] = []
    for src in [
        msg.get("best_oa_location") or {},
        msg.get("primary_location") or {},
    ]:
        u = src.get("pdf_url")
        if u and u not in out:
            out.append(u)
    u = (msg.get("open_access") or {}).get("oa_url")
    if u and u not in out:
        out.append(u)
    for loc in msg.get("locations") or []:
        u = loc.get("pdf_url")
        if u and u not in out:
            out.append(u)
    _lookup_urls_cache[ck] = tuple(out)
    return out


def core_pdf_urls(doi: str) -> list[str]:
    ck_cache = ("core", doi.lower().strip())
    if ck_cache in _lookup_urls_cache:
        return list(_lookup_urls_cache[ck_cache])
    q = urllib.parse.quote(f'doi:"{doi}"', safe="")
    url = f"https://api.core.ac.uk/v3/search/works?q={q}&limit=1"
    xh: dict[str, str] = {}
    ck = os.environ.get("AI4RESEARCH_CORE_API_KEY", "").strip()
    if ck:
        xh["Authorization"] = f"Bearer {ck}"
    try:
        _rate_limit_source("core")
        data, _ = http_get_urllib(url, extra_headers=xh or None)
        msg = json.loads(data.decode("utf-8", errors="replace"))
    except Exception:
        _lookup_urls_cache[ck_cache] = ()
        return []
    out: list[str] = []
    for result in (msg.get("results") or [])[:1]:
        u = result.get("downloadUrl")
        if u and u not in out:
            out.append(u)
    _lookup_urls_cache[ck_cache] = tuple(out)
    return out


def _europepmc_parse_fulltext_pdfs(r0: dict) -> list[str]:
    ftl = r0.get("fullTextUrlList")
    if not ftl:
        return []
    items = ftl.get("fullTextUrl")
    if isinstance(items, dict):
        items = [items]
    elif not items:
        return []
    out: list[str] = []
    for item in items:
        if item.get("documentStyle") != "pdf":
            continue
        if item.get("availabilityCode") != "F":
            continue
        u = item.get("url")
        if u and "pdf" in u.lower():
            out.append(u)
    return out


def europepmc_pdf_urls(doi: str) -> list[str]:
    ck = ("europepmc_doi", doi.lower().strip())
    if ck in _lookup_urls_cache:
        return list(_lookup_urls_cache[ck])
    q = urllib.parse.quote(doi, safe="")
    url = f"https://www.ebi.ac.uk/europepmc/webservices/rest/search?query=DOI:{q}&resultType=core&format=json"
    try:
        _rate_limit_source("europepmc")
        data, _ = http_get_urllib(url)
        results = json.loads(data.decode("utf-8", errors="replace"))["resultList"].get(
            "result"
        ) or []
        if not results:
            _lookup_urls_cache[ck] = ()
            return []
        out = _europepmc_parse_fulltext_pdfs(results[0])
        _lookup_urls_cache[ck] = tuple(out)
        return out
    except Exception:
        _lookup_urls_cache[ck] = ()
        return []


def europepmc_pdf_urls_for_pmc_id(pmc_numeric: str) -> list[str]:
    """Concrete PDF URLs via Europe PMC metadata (NIHMS-named paths), not directory stubs."""
    pid = pmc_numeric.strip().replace("PMC", "")
    if not pid.isdigit():
        return []
    ck = ("europepmc_pmcid", pid)
    if ck in _lookup_urls_cache:
        return list(_lookup_urls_cache[ck])
    q = urllib.parse.quote(f"EXT_ID:PMC{pid}", safe="")
    url = (
        "https://www.ebi.ac.uk/europepmc/webservices/rest/search?"
        f"query={q}&resultType=core&format=json"
    )
    try:
        _rate_limit_source("europepmc")
        data, _ = http_get_urllib(url)
        results = json.loads(data.decode("utf-8", errors="replace"))["resultList"].get(
            "result"
        ) or []
        if not results:
            _lookup_urls_cache[ck] = ()
            return []
        out = _europepmc_parse_fulltext_pdfs(results[0])
        _lookup_urls_cache[ck] = tuple(out)
        return out
    except Exception:
        _lookup_urls_cache[ck] = ()
        return []


def crossref_pdf_urls(doi: str) -> list[str]:
    ck = ("crossref", doi.lower().strip())
    if ck in _lookup_urls_cache:
        return list(_lookup_urls_cache[ck])
    q = urllib.parse.quote(doi, safe="")
    url = f"https://api.crossref.org/works/{q}"
    try:
        _rate_limit_source("crossref")
        data, _ = http_get_urllib(url)
        msg = json.loads(data.decode("utf-8", errors="replace"))["message"]
    except Exception:
        _lookup_urls_cache[ck] = ()
        return []
    out: list[str] = []
    for link in msg.get("link") or []:
        u = link.get("URL") or ""
        ct = (link.get("content-type") or "").lower()
        if not u:
            continue
        if ct == "application/pdf" or "/pdf" in u.lower() or u.lower().endswith(".pdf"):
            if u not in out:
                out.append(u)
    _lookup_urls_cache[ck] = tuple(out)
    return out


def _reg_domain(netloc: str) -> str:
    host = (netloc or "").split(":")[0].lower()
    parts = host.split(".")
    if len(parts) >= 2:
        return ".".join(parts[-2:])
    return host


def _resolve_url(u: str, base_url: str) -> str:
    u = html_module.unescape(u.strip())
    if u.startswith("//"):
        return "https:" + u
    if u.startswith("/"):
        p = urllib.parse.urlparse(base_url)
        return f"{p.scheme}://{p.netloc}{u}"
    return u


def extract_pdf_urls_from_html(body: bytes, base_url: str) -> list[str]:
    """Prefer citation_pdf_url meta and <link type=application/pdf>; also catches
    /doi/pdf/ publisher links; only accepts same-registrable-domain hrefs for generic .pdf links."""
    text = body.decode("utf-8", errors="replace")
    meta_urls: list[str] = []
    for pat in (
        r'citation_pdf_url"\s+content="([^"]+)"',
        r"citation_pdf_url'\s+content='([^']+)'",
        r'<meta\s+name="citation_pdf_url"\s+content="([^"]+)"',
        # <link rel="alternate" type="application/pdf" href="...">
        r'<link[^>]+type=["\']application/pdf["\'][^>]+href=["\']([^"\']+)["\']',
        r'<link[^>]+href=["\']([^"\']+)["\'][^>]+type=["\']application/pdf["\']',
    ):
        for m in re.finditer(pat, text, re.I):
            meta_urls.append(m.group(1))

    base_dom = _reg_domain(urllib.parse.urlparse(base_url).netloc)
    href_urls: list[str] = []

    # Same-domain .pdf hrefs
    for m in re.finditer(r'href="([^"]+\.pdf[^"]*)"', text, re.I):
        u = m.group(1)
        if "javascript" in u.lower():
            continue
        abs_u = _resolve_url(u, base_url)
        if not abs_u.startswith("http"):
            continue
        if _reg_domain(urllib.parse.urlparse(abs_u).netloc) != base_dom:
            continue
        href_urls.append(abs_u)

    # Publisher /doi/pdf/, /doi/epdf/, /doi/pdfplus/ patterns (Taylor & Francis, Oxford, Wiley, etc.)
    for m in re.finditer(r'href="([^"]*?/doi/(?:pdf|epdf|pdfplus)/[^"]*)"', text, re.I):
        u = m.group(1)
        if "javascript" in u.lower():
            continue
        abs_u = _resolve_url(u, base_url)
        if not abs_u.startswith("http"):
            continue
        if "/epdf/" in abs_u:
            if "onlinelibrary.wiley.com" in abs_u:
                abs_u = abs_u.replace("/epdf/", "/pdfdirect/") + "?download=true"
            elif "tandfonline.com" in abs_u:
                abs_u = abs_u.replace("/epdf/", "/pdf/") + "?download=true"
            else:
                abs_u = abs_u.replace("/epdf/", "/pdf/")

        if abs_u not in href_urls:
            href_urls.append(abs_u)

    ordered = meta_urls + href_urls
    abs_urls: list[str] = []
    for u in ordered:
        u = _resolve_url(u, base_url)
        if not u.startswith("http"):
            continue
        if u not in abs_urls:
            abs_urls.append(u)
    return abs_urls


def _curl_cffi_fetch_pdf(url: str, referer: str) -> bytes | None:
    try:
        from curl_cffi import requests as creq
    except ImportError:
        return None
    try:
        r = creq.get(
            url,
            impersonate="chrome136",
            headers={
                "User-Agent": USER_AGENT,
                "Referer": referer,
                "Accept": "application/pdf,*/*",
            },
            timeout=120,
        )
        body = getattr(r, "content", None) or b""
        ok, _why = _validate_pdf_bytes(body)
        if r.status_code == 200 and ok:
            return body
    except Exception:
        pass
    return None


def _playwright_request_pdf(url: str, referer: str) -> bytes | None:
    if os.environ.get("AI4RESEARCH_PLAYWRIGHT", "").strip().lower() not in (
        "1",
        "true",
        "yes",
    ):
        return None
    try:
        from playwright.sync_api import sync_playwright  # type: ignore[import-untyped]
    except ImportError:
        print(
            "AI4RESEARCH_PLAYWRIGHT=1 requires: pip install playwright && playwright install chromium",
            file=sys.stderr,
        )
        return None
    hdr = {
        "User-Agent": USER_AGENT,
        "Referer": referer,
        "Accept": "application/pdf,*/*",
    }
    try:
        with sync_playwright() as p:
            ctx = p.request.new_context(extra_http_headers=hdr)
            resp = ctx.get(url, timeout=120_000)
            if getattr(resp, "ok", resp.status == 200):
                body = resp.body()
                ok, _why = _validate_pdf_bytes(body)
                if ok:
                    return body
    except Exception:
        pass
    return None


def fetch_pdf_bytes(
    url: str,
    referer: str | None = None,
    doi_landing: str | None = None,
) -> bytes | None:
    """GET URL with redirects; succeed when body begins with %PDF."""
    ref = doi_landing or referer or referer_for_url(url)

    # curl_cffi impersonates Chrome TLS — best first attempt for publisher pages.
    cf = _curl_cffi_fetch_pdf(url, ref)
    if cf:
        return cf

    backoff = _RETRYABLE_STATUS | {403}

    for attempt in range(_MAX_RETRIES):
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": USER_AGENT,
                "Accept": "application/pdf,*/*;q=0.9",
                "Accept-Language": "en-US,en;q=0.9",
                "Accept-Encoding": "gzip, deflate",
                "Referer": ref,
                "Sec-Fetch-Site": "cross-site",
                "Sec-Fetch-Mode": "navigate",
                "Sec-Fetch-Dest": "document",
                "Sec-Fetch-User": "?1",
                "Sec-Ch-Ua": '"Chromium";v="136", "Google Chrome";v="136", "Not.A/Brand";v="99"',
                "Sec-Ch-Ua-Mobile": "?0",
                "Sec-Ch-Ua-Platform": '"Windows"',
            },
            method="GET",
        )
        try:
            with _urlopen_with_cookies(req, 120) as resp:
                data = _decompress_if_needed(resp.read(), resp)
            ok, _why = _validate_pdf_bytes(data)
            if ok:
                return data
        except urllib.error.HTTPError as e:
            try:
                e.read()
            except Exception:
                pass
            if (
                getattr(e, "code", None) in backoff
                and attempt < _MAX_RETRIES - 1
            ):
                time.sleep((2**attempt) * 0.35 + random.random() * 0.22)
                continue
            break
        except (urllib.error.URLError, TimeoutError, OSError):
            if attempt < _MAX_RETRIES - 1:
                time.sleep((2**attempt) * 0.25 + random.random() * 0.15)
                continue
            break

    c = curl_fetch(url, referer=ref)
    if c:
        ok, _why = _validate_pdf_bytes(c)
        if ok:
            return c
    return _playwright_request_pdf(url, ref)


def try_urls(
    urls: list[str],
    *,
    doi_landing: str | None = None,
) -> tuple[bytes | None, str | None]:
    for u in urls:
        data = fetch_pdf_bytes(u, doi_landing=doi_landing)
        if data:
            return data, u
        time.sleep(0.2)
    return None, None


def fetch_landing_html(doi: str) -> tuple[bytes, str] | tuple[None, str]:
    landing = f"https://doi.org/{urllib.parse.quote(doi, safe='')}"
    try:
        body, final = http_get_urllib(landing)
        return body, final
    except Exception:
        c = curl_fetch(landing, referer="https://doi.org/")
        if c:
            return c, landing
        return None, f"failed: {landing}"


def synthesize_publisher_pdf_urls(doi: str) -> list[str]:
    """Guess direct PDF URLs based on DOI prefix to bypass HTML landing pages."""
    urls = []
    if doi.startswith("10.1080/"):
        urls.append(f"https://www.tandfonline.com/doi/pdf/{doi}?download=true")
    elif doi.startswith("10.1177/"):
        urls.append(f"https://journals.sagepub.com/doi/pdf/{doi}?download=true")
    elif doi.startswith("10.1038/"):
        urls.append(f"https://www.nature.com/articles/{doi.split('/')[-1]}.pdf")
    elif doi.startswith("10.1111/") or doi.startswith("10.1002/"):
        urls.append(f"https://onlinelibrary.wiley.com/doi/pdfdirect/{doi}?download=true")
    return urls


def download_for_doi(doi: str, use_scihub: bool = False) -> tuple[bool, bytes, str]:
    landing = f"https://doi.org/{urllib.parse.quote(doi, safe='')}"
    uw = unpaywall_lookup(doi)
    if uw:
        ulist = iter_unpaywall_pdf_urls(uw)
        data, src = try_urls(ulist, doi_landing=landing)
        if data:
            return True, data, f"unpaywall: {src}"
        for u in unpaywall_landing_follow_urls(uw, ulist):
            got = fetch_pdf_bytes(u, doi_landing=landing)
            if got:
                return True, got, f"unpaywall-follow: {u}"

    ss = semantic_scholar_lookup(doi)
    if ss:
        sslist = iter_semantic_scholar_pdf_urls(ss)
        data, src = try_urls(sslist, doi_landing=landing)
        if data:
            return True, data, f"semanticscholar: {src}"

    oalist = openalex_pdf_urls(doi)
    data, src = try_urls(oalist, doi_landing=landing)
    if data:
        return True, data, f"openalex: {src}"

    xlist = crossref_pdf_urls(doi)
    data, src = try_urls(xlist, doi_landing=landing)
    if data:
        return True, data, f"crossref: {src}"

    emlist = europepmc_pdf_urls(doi)
    data, src = try_urls(emlist, doi_landing=landing)
    if data:
        return True, data, f"europepmc: {src}"

    clist = core_pdf_urls(doi)
    data, src = try_urls(clist, doi_landing=landing)
    if data:
        return True, data, f"core: {src}"

    if ss:
        plist = preprint_pdf_urls(ss)
        data, src = try_urls(plist, doi_landing=landing)
        if data:
            return True, data, f"preprint: {src}"

    synth_urls = synthesize_publisher_pdf_urls(doi)
    data, src = try_urls(synth_urls, doi_landing=landing)
    if data:
        return True, data, f"synthesized: {src}"

    result = fetch_landing_html(doi)
    if result[0] is None:
        return False, b"", str(result[1])
    html_bytes, final = result

    ok_direct, why_direct = _validate_pdf_bytes(html_bytes)
    if ok_direct:
        return True, html_bytes, f"doi redirect PDF: {final}"
    if html_bytes[:4] == b"%PDF":
        return False, b"", f"doi redirect invalid PDF: {why_direct}"

    for candidate in extract_pdf_urls_from_html(html_bytes, final):
        data = fetch_pdf_bytes(candidate, doi_landing=landing)
        if data:
            return True, data, f"html: {candidate}"
        time.sleep(0.2)

    if use_scihub:
        for sh_domain in ("sci-hub.se", "sci-hub.st", "sci-hub.ru"):
            sh_url = f"https://{sh_domain}/{doi}"
            try:
                sh_body, _ = http_get_urllib(sh_url)
                sh_text = sh_body.decode("utf-8", errors="replace")
                m = re.search(r'<iframe\s+src="([^"]+)"\s+id="pdf"', sh_text, re.I)
                if m:
                    pdf_u = m.group(1)
                    if pdf_u.startswith("//"):
                        pdf_u = "https:" + pdf_u
                    elif pdf_u.startswith("/"):
                        pdf_u = f"https://{sh_domain}" + pdf_u
                    sh_data = fetch_pdf_bytes(pdf_u, doi_landing=sh_url)
                    if sh_data:
                        return True, sh_data, f"sci-hub: {pdf_u}"
            except Exception:
                pass

    extra = ""
    if uw:
        extra = f" is_oa={uw.get('is_oa')} title={uw.get('title', '')[:60]!r}"
    return False, b"", f"no PDF found{extra}"


def run_xlsx_citation_workflow(
    xlsx_path: Path,
    *,
    out_dir: Path,
    citation_column: str,
    context_column_names: tuple[str, ...],
    extra_boost_tokens: frozenset[str],
    use_scholarly: bool,
    resolve_only: bool,
    summary_md_path: Path,
    search_urls_on_fail: bool = False,
    web_html_resolve: bool = False,
    use_scihub: bool = False,
) -> None:
    xp = xlsx_path.expanduser().resolve()
    if not xp.is_file():
        print(f"XLSX not found: {xp}", file=sys.stderr)
        raise SystemExit(1)
    out_dir = out_dir.expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    summary_md_path = summary_md_path.expanduser().resolve()
    summary_md_path.parent.mkdir(parents=True, exist_ok=True)

    pairs, warns = iter_xlsx_author_date_citations(
        xp, citation_column, context_column_names
    )
    for w in warns:
        print(w, file=sys.stderr)

    by_sheet: dict[str, list[list[str]]] = {}
    for sheet_name, _r_i, _cite, _ctx in pairs:
        by_sheet.setdefault(sheet_name, [])

    print(f"\n=== XLSX {xp.name} ({len(pairs)} citation cells) ===")

    completed_all_rows = False

    def _write_summary_md() -> None:
        n_done = sum(len(rows) for rows in by_sheet.values())
        sheet_tables = [(sn, by_sheet[sn]) for sn in by_sheet if by_sheet[sn]]
        ctx_note = (
            f"  Context columns: {', '.join(context_column_names)}\n"
            if context_column_names
            else ""
        )
        kw_note = (
            f"  Extra keyword boost tokens: {len(extra_boost_tokens)}\n"
            if extra_boost_tokens
            else ""
        )
        sch_note = f"  Scholarly (Google Scholar) fallback: {use_scholarly}\n"
        url_note = (
            "  On unresolved / ambiguous (>5): web search URLs printed to stderr (--search-urls)\n"
            if search_urls_on_fail
            else ""
        )
        ddg_note = (
            f"  DuckDuckGo Lite HTML fallback: {web_html_resolve}\n"
            if web_html_resolve
            else ""
        )
        partial_note = (
            ""
            if completed_all_rows
            else (
                f"\n  **Partial run:** only {n_done} summary row(s) "
                f"from {len(pairs)} citation cell(s) "
                "were processed (e.g. interrupted or error). Tables below reflect "
                "that subset.\n"
            )
        )
        intro = (
            f"\nCitation cells in spreadsheet: {len(pairs)}  "
            f"Summary table rows: {n_done}  Resolve-only: {resolve_only}\n"
            f"{ctx_note}{kw_note}{sch_note}{url_note}{ddg_note}{partial_note}"
        )
        md = build_xlsx_download_markdown(
            xlsx_path=xp, sheet_tables=sheet_tables, intro_extra=intro
        )
        summary_md_path.write_text(md, encoding="utf-8")
        summary_csv_path = summary_md_path.with_suffix(".csv")
        write_xlsx_download_csv(csv_path=summary_csv_path, sheet_tables=sheet_tables)
        state = "complete" if completed_all_rows else "partial"
        print(
            f"\nWrote {summary_md_path} and {summary_csv_path} "
            f"({state}, {n_done} row(s))"
        )

    try:
        for sheet_name, _r_i, cite_text, ctx_blob in pairs:
            title_cell = "-"
            doi_cell = ""
            src_cell = ""
            status = ""
            detail = ""
            authors_cell = ""
            year_cell = ""
            journal_cell = ""
            parsed = parse_author_year_citation(cite_text)
            if parsed is None:
                status = "unparsed"
                detail = "does not match author (year) pattern"
                row = [
                    cite_text,
                    authors_cell,
                    year_cell,
                    journal_cell,
                    title_cell,
                    doi_cell,
                    status,
                    detail,
                    src_cell,
                ]
                by_sheet[sheet_name].append(row)
                print(
                    f"  [{sheet_name}] {cite_text[:56]!r} -> {status} {doi_cell or '-'}"
                )
            else:
                row_boost = title_boost_tokens(ctx_blob) | extra_boost_tokens
                boost_for = row_boost if row_boost else None
                res = resolve_citation_to_work(
                    parsed,
                    boost_for,
                    use_scholarly=use_scholarly,
                    keyword_boost_required=extra_boost_tokens
                    if extra_boost_tokens
                    else None,
                    try_web_html_resolve=web_html_resolve,
                )
                ac = res.ambiguous_candidates
                if ac is not None and 2 <= len(ac) <= 5:
                    print(
                        f"  [{sheet_name}] {cite_text[:56]!r} -> ambiguous "
                        f"({len(ac)} candidates)",
                    )
                    for cand in ac:
                        by_sheet[sheet_name].append(
                            [
                                cite_text,
                                cand.authors or "",
                                (
                                    str(cand.publication_year)
                                    if cand.publication_year is not None
                                    else ""
                                ),
                                cand.journal or "",
                                cand.title if cand.title else "-",
                                cand.doi or "",
                                "ambiguous",
                                cand.detail,
                                cand.source,
                            ]
                        )
                    if search_urls_on_fail:
                        print_citation_web_search_urls(
                            parsed,
                            boost_for,
                            cite_label=f"[{sheet_name}] {cite_text[:72]!r}",
                        )
                else:
                    title_cell = res.title if res.title else "-"
                    doi_cell = res.doi or ""
                    src_cell = res.source
                    authors_cell = res.authors or ""
                    year_cell = (
                        str(res.publication_year)
                        if res.publication_year is not None
                        else ""
                    )
                    journal_cell = res.journal or ""
                    if res.status == "ambiguous (>5)":
                        status = "ambiguous (>5)"
                        detail = res.detail
                    elif res.status == "ambiguous":
                        status = "ambiguous"
                        detail = res.detail
                    elif res.status != "resolved" or not res.doi:
                        status = "unresolved"
                        detail = res.detail
                    elif resolve_only:
                        status = "resolved"
                        detail = f"{res.detail}; resolve-only (no PDF)"
                    else:
                        fn = slug_from_doi(res.doi) + ".pdf"
                        dest = out_dir / fn
                        if dest.exists() and dest.stat().st_size > 1000:
                            status = "cached"
                            detail = "pdf already on disk"
                        else:
                            ok, data, msg = download_for_doi(res.doi, use_scihub=use_scihub)
                            if ok and data:
                                dest.write_bytes(data)
                                status = "ok"
                                detail = msg
                            else:
                                status = "fail"
                                detail = msg

                    row = [
                        cite_text,
                        authors_cell,
                        year_cell,
                        journal_cell,
                        title_cell,
                        doi_cell,
                        status,
                        detail,
                        src_cell,
                    ]
                    by_sheet[sheet_name].append(row)
                    print(
                        f"  [{sheet_name}] {cite_text[:56]!r} -> {status} {doi_cell or '-'}"
                    )
                    if search_urls_on_fail and status in (
                        "unresolved",
                        "ambiguous (>5)",
                    ):
                        print_citation_web_search_urls(
                            parsed,
                            boost_for,
                            cite_label=f"[{sheet_name}] {cite_text[:72]!r}",
                        )

            if parsed is not None:
                time.sleep(0.35)
            time.sleep(0.05)
        completed_all_rows = True
    except KeyboardInterrupt:
        print(
            "\nInterrupted — writing summary with rows processed so far...",
            file=sys.stderr,
        )
        raise
    finally:
        _write_summary_md()


def _parse_cli(argv: list[str] | None) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Download PDFs cited in chatbots/*.md, XLSX author-year column, one DOI, or one citation.",
    )
    p.add_argument(
        "--cite",
        metavar="AUTHOR_YEAR",
        default=None,
        help=(
            "Single author–year citation (must end with ``(YYYY)``), e.g. ``Smith, Jones (2020)``. "
            "Resolve DOI like XLSX mode; use --resolve-only to print result without downloading a PDF."
        ),
    )
    p.add_argument(
        "--doi",
        metavar="DOI",
        help="Resolve and download one DOI to --out (default: pdfs/single/<slug>.pdf)",
    )
    p.add_argument(
        "--out",
        type=Path,
        help="Output path when using --doi or --cite (default pdfs/single/<slug>.pdf)",
    )
    p.add_argument(
        "--dry-run",
        action="store_true",
        help="With --doi: resolve only; print result, do not write",
    )
    p.add_argument(
        "--xlsx",
        type=Path,
        metavar="PATH",
        help="Read citation column from XLSX; resolve DOIs and download PDFs",
    )
    p.add_argument(
        "--xlsx-out-dir",
        type=Path,
        metavar="DIR",
        help="PDF output directory (default: pdfs/<xlsx_stem>/)",
    )
    p.add_argument(
        "--citation-column",
        default="Author (date)",
        metavar="NAME",
        help='Header text to locate the citation column (default: "Author (date)")',
    )
    p.add_argument(
        "--context-column",
        action="append",
        default=None,
        metavar="NAME",
        help=(
            "Optional XLSX column header (same header row as citation column). "
            "Cell text on each row is tokenized; tokens boost Crossref/OpenAlex ranking "
            "via title overlap (not mixed into Crossref/OpenAlex query strings). Repeat "
            "flag for multiple columns."
        ),
    )
    p.add_argument(
        "--keyword-boost",
        default="",
        metavar="WORDS",
        help=(
            "Comma- or space-separated topic words (e.g. childhood,poverty,income) "
            "merged into every row's keyword set (and into --cite resolution). Used to score/rank Crossref and OpenAlex "
            "hits by title overlap (not added to those APIs' search strings, which would "
            "often hide the correct paper). Appended to the Scholarly query and used as "
            "extra words for DuckDuckGo Lite / --search-urls queries. With "
            "--scholarly, Scholar results must include every "
            "keyword-boost term in the title."
        ),
    )
    p.add_argument(
        "--scholarly",
        action="store_true",
        help=(
            "With --xlsx / --cite: if Crossref does not resolve a citation, try Google Scholar "
            "via the optional scholarly package (pip install scholarly). May be "
            "rate-limited; see AI4RESEARCH_SCHOLARLY=1 to enable without this flag."
        ),
    )
    p.add_argument(
        "--search-urls",
        action="store_true",
        help=(
            "With --xlsx / --cite: print Google, Scholar, and DuckDuckGo search URLs (stderr) "
            "when resolution fails—same query string as Scholarly (authors + year + "
            "--keyword-boost tokens)."
        ),
    )
    p.add_argument(
        "--web-html-resolve",
        action="store_true",
        help=(
            "After Crossref and OpenAlex (and optional Scholarly), fetch DuckDuckGo Lite HTML "
            "and take the first DOI whose Crossref metadata passes first-author match. "
            "Also AI4RESEARCH_WEB_HTML_RESOLVE=1."
        ),
    )
    p.add_argument(
        "--resolve-only",
        action="store_true",
        help="With --xlsx / --cite: resolve only; do not download PDFs (cite prints summary to stdout)",
    )
    p.add_argument(
        "--sci-hub",
        action="store_true",
        help="Use Sci-Hub as a last resort for PDF downloads if Open Access and publisher HTML scraping fail.",
    )
    return p.parse_args(argv if argv is not None else sys.argv[1:])


def main(argv: list[str] | None = None) -> None:
    args = _parse_cli(argv)

    cite_present = bool((args.cite or "").strip())
    _modes = sum(bool(x) for x in (args.doi, args.xlsx, cite_present))
    if _modes > 1:
        print("Use only one of --doi, --xlsx, or --cite.", file=sys.stderr)
        raise SystemExit(2)

    if args.xlsx:
        PDFS.mkdir(parents=True, exist_ok=True)
        xp = Path(args.xlsx)
        stem = xp.stem
        out_dir = Path(args.xlsx_out_dir) if args.xlsx_out_dir else PDFS / stem
        summary_md = PDFS / f"{stem}_download_summary.md"
        ctx_cols = tuple(
            c.strip()
            for c in (args.context_column or [])
            if c and str(c).strip()
        )
        extra_kw = title_boost_tokens(
            (args.keyword_boost or "").replace(",", " ")
        )
        use_scholarly = bool(args.scholarly) or os.environ.get(
            "AI4RESEARCH_SCHOLARLY", ""
        ).strip().lower() in ("1", "true", "yes")
        web_html = bool(args.web_html_resolve) or _env_truthy(
            "AI4RESEARCH_WEB_HTML_RESOLVE"
        )
        run_xlsx_citation_workflow(
            xp,
            out_dir=out_dir,
            citation_column=args.citation_column,
            context_column_names=ctx_cols,
            extra_boost_tokens=extra_kw,
            use_scholarly=use_scholarly,
            resolve_only=bool(args.resolve_only),
            summary_md_path=summary_md,
            search_urls_on_fail=bool(args.search_urls),
            web_html_resolve=web_html,
            use_scihub=bool(args.sci_hub),
        )
        return

    cite_raw = (args.cite or "").strip()
    if cite_raw:
        if args.dry_run:
            print("--dry-run applies only with --doi.", file=sys.stderr)
            raise SystemExit(2)
        PDFS.mkdir(parents=True, exist_ok=True)
        extra_kw = title_boost_tokens(
            (args.keyword_boost or "").replace(",", " ")
        )
        use_scholarly = bool(args.scholarly) or os.environ.get(
            "AI4RESEARCH_SCHOLARLY", ""
        ).strip().lower() in ("1", "true", "yes")
        web_html = bool(args.web_html_resolve) or _env_truthy(
            "AI4RESEARCH_WEB_HTML_RESOLVE"
        )
        parsed = parse_author_year_citation(cite_raw)
        if parsed is None:
            print(
                f"Citation does not match author (year) tail pattern: {cite_raw!r}",
                file=sys.stderr,
            )
            raise SystemExit(2)
        boost_for = extra_kw if extra_kw else None
        res = resolve_citation_to_work(
            parsed,
            boost_for,
            use_scholarly=use_scholarly,
            keyword_boost_required=extra_kw if extra_kw else None,
            try_web_html_resolve=web_html,
        )
        if res.status == "ambiguous (>5)":
            if args.search_urls:
                print_citation_web_search_urls(
                    parsed, boost_for, cite_label=cite_raw[:88]
                )
            print(
                f"ambiguous (>5) {cite_raw!r}: {res.detail}",
                file=sys.stderr,
            )
            raise SystemExit(1)
        ac = res.ambiguous_candidates
        if ac is not None and 2 <= len(ac) <= 5:
            print(f"ambiguous {cite_raw!r}: {res.detail}")
            for cand in ac:
                print(
                    f"  candidate: doi={cand.doi} title={(cand.title or '-')!r} "
                    f"year={cand.publication_year} journal={cand.journal or '-'} "
                    f"({cand.detail})"
                )
            if args.search_urls:
                print_citation_web_search_urls(
                    parsed, boost_for, cite_label=cite_raw[:88]
                )
            raise SystemExit(1)
        if res.status == "ambiguous":
            if args.search_urls:
                print_citation_web_search_urls(
                    parsed, boost_for, cite_label=cite_raw[:88]
                )
            print(
                f"ambiguous {cite_raw!r}: {res.detail}\n"
                f"  title={res.title!r}\n"
                f"  source={res.source!r}",
                file=sys.stderr,
            )
            raise SystemExit(1)
        if res.status != "resolved" or not res.doi:
            if args.search_urls:
                print_citation_web_search_urls(
                    parsed, boost_for, cite_label=cite_raw[:88]
                )
            print(
                f"unresolved {cite_raw!r}: {res.detail}\n"
                f"  title={res.title!r}\n"
                f"  source={res.source!r}",
                file=sys.stderr,
            )
            raise SystemExit(1)
        doi = res.doi.strip()
        print(
            f"resolved {cite_raw!r}\n"
            f"  doi={doi}\n"
            f"  title={res.title or '-'}\n"
            f"  source={res.source}"
        )
        if args.resolve_only:
            return
        ok, data, msg = download_for_doi(doi, use_scihub=bool(args.sci_hub))
        dest = (
            Path(args.out)
            if args.out
            else (PDFS / "single" / (slug_from_doi(doi) + ".pdf"))
        )
        dest.parent.mkdir(parents=True, exist_ok=True)
        if ok and data:
            dest.write_bytes(data)
            print(f"OK {doi} -> {dest} ({msg})")
            return
        print(f"FAIL {doi}: {msg}", file=sys.stderr)
        raise SystemExit(1)

    if args.doi:
        doi = args.doi.strip()
        PDFS.mkdir(parents=True, exist_ok=True)
        ok, data, msg = download_for_doi(doi, use_scihub=bool(args.sci_hub))
        if args.dry_run:
            print(
                f"dry-run {doi!r}: ok={ok} bytes={len(data) if ok else 0} detail={msg}"
            )
            return
        dest = (
            Path(args.out)
            if args.out
            else (PDFS / "single" / (slug_from_doi(doi) + ".pdf"))
        )
        dest.parent.mkdir(parents=True, exist_ok=True)
        if ok and data:
            dest.write_bytes(data)
            print(f"OK {doi} -> {dest} ({msg})")
            return
        print(f"FAIL {doi}: {msg}", file=sys.stderr)
        raise SystemExit(1)

    mapping = [
        ("methods_consultant_psychology.md", "methods_consultant_psychology"),
        ("methods_consultant_econometrics.md", "methods_consultant_econometrics"),
        ("methods_consultant_epidemiology.md", "methods_consultant_epidemiology"),
        ("methods_consultant.md", "methods_consultant"),
    ]

    PDFS.mkdir(parents=True, exist_ok=True)
    summary: list[tuple[str, str, str, str]] = []

    for md_name, folder in mapping:
        md_path = CHATBOTS / md_name
        if not md_path.exists():
            print(f"Skip missing {md_path}")
            continue
        text = md_path.read_text(encoding="utf-8")
        dois = find_dois_in_md(text)
        out_dir = PDFS / folder
        out_dir.mkdir(parents=True, exist_ok=True)
        print(f"\n=== {folder} ({len(dois)} DOIs) ===")

        for doi in dois:
            fn = slug_from_doi(doi) + ".pdf"
            dest = out_dir / fn
            if dest.exists() and dest.stat().st_size > 1000:
                print(f"  OK (cached): {doi}")
                summary.append((folder, doi, "cached", "-"))
                continue

            ok, data, msg = download_for_doi(doi, use_scihub=bool(args.sci_hub))

            if ok and data:
                dest.write_bytes(data)
                print(f"  OK: {doi} -> {fn} ({msg})")
                summary.append((folder, doi, "ok", msg))
            else:
                print(f"  FAIL: {doi} ({msg})")
                summary.append((folder, doi, "fail", msg))

            time.sleep(0.4)

    fail_log = PDFS / "_download_summary.tsv"
    hdr = "folder\tdoi\tstatus\tdetail"
    body_lines = ["\t".join(x) for x in summary]
    fail_log.write_text(hdr + "\n" + "\n".join(body_lines) + ("\n" if body_lines else ""), encoding="utf-8")
    print(f"\nWrote {fail_log}")


if __name__ == "__main__":
    main()
